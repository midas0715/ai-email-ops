import io
import os

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class ExtractionError(Exception):
    """The file could not be turned into text. The message is safe to show to the user."""


def extract_text(filename: str, data: bytes) -> str:
    extension = os.path.splitext(filename.lower())[1]
    if extension == ".txt":
        text = data.decode("utf-8-sig", errors="replace")
    elif extension == ".pdf":
        text = _pdf_text(data)
    elif extension == ".docx":
        text = _docx_text(data)
    else:
        raise ExtractionError("Unsupported file type")

    text = text.replace("\x00", "").strip()  # Postgres cannot store NUL characters
    if not text:
        raise ExtractionError(
            "No readable text found. Scanned PDFs and images need OCR, which is not supported yet."
        )
    return text


def _pdf_text(data: bytes) -> str:
    if not data.startswith(b"%PDF"):
        raise ExtractionError("This file is not a real PDF")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("Password-protected PDFs are not supported")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:  # damaged files can raise many different errors
        raise ExtractionError("The PDF could not be read") from exc


def _docx_text(data: bytes) -> str:
    if not data.startswith(b"PK"):  # a .docx is a zip file
        raise ExtractionError("This file is not a real DOCX")
    try:
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise ExtractionError("The DOCX could not be read") from exc